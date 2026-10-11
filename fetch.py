#!/usr/bin/env python3
"""Fetch free, delayed market data into data.json. No API keys. Each source fails independently,
keeping its last good value."""
import json, re, io, datetime as dt, html, email.utils, urllib.parse
import requests

UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0 Safari/537.36"}
NOW = dt.datetime.now(dt.timezone.utc)
ISO = lambda d=None: (d or NOW).strftime("%Y-%m-%dT%H:%M:%SZ")
LB_PER_T = 2204.62262
try:
    OLD = json.load(open("data.json"))
except Exception:
    OLD = {}
OUT = {"generated": ISO(), "errors": {}}

def get(url, **kw):
    r = requests.get(url, headers=UA, timeout=25, **kw); r.raise_for_status(); return r

def section(name, fn):
    try:
        v = fn(); v.setdefault("asof", ISO()); v["ok"] = True; OUT[name] = v
    except Exception as e:
        OUT["errors"][name] = f"{type(e).__name__}: {e}"[:200]
        prev = OLD.get(name)
        OUT[name] = dict(prev, ok=False) if prev else {"ok": False, "unavailable": True}

def yahoo(sym, rng="1y"):
    try:
        j = get(f"https://query1.finance.yahoo.com/v8/finance/chart/{sym}?range={rng}&interval=1d").json()
        res = j["chart"]["result"][0]; m = res["meta"]
        pts = [(t, c) for t, c in zip(res["timestamp"], res["indicators"]["quote"][0]["close"]) if c is not None]
        price = m.get("regularMarketPrice") or pts[-1][1]
        asof = ISO(dt.datetime.fromtimestamp(m.get("regularMarketTime", pts[-1][0]), dt.timezone.utc))
        cur = m.get("currency")
    except Exception:  # Stooq fallback (daily CSV)
        s = sym.lower().replace("=f", ".f").replace(".to", ".ca").replace(".ax", ".au").replace(".l", ".uk")
        if "." not in s: s += ".us"
        rows = get(f"https://stooq.com/q/d/l/?s={s}&i=d").text.strip().splitlines()[1:]
        pts = [(int(dt.datetime.strptime(r.split(",")[0], "%Y-%m-%d").replace(tzinfo=dt.timezone.utc).timestamp()), float(r.split(",")[4])) for r in rows[-260:]]
        price = pts[-1][1]; asof = ISO(dt.datetime.fromtimestamp(pts[-1][0], dt.timezone.utc)); cur = None
    # previous close = last daily close before today's bar
    prev = pts[-2][1] if len(pts) > 1 else price
    def back(days):
        cut = NOW.timestamp() - days * 86400
        c = [p for t, p in pts if t <= cut]; return c[-1] if c else pts[0][1]
    ystart = dt.datetime(NOW.year, 1, 1, tzinfo=dt.timezone.utc).timestamp()
    ytd_base = ([p for t, p in pts if t < ystart] or [pts[0][1]])[-1]
    pct = lambda a, b: round((a / b - 1) * 100, 2) if b else None
    return {"price": price, "currency": cur, "chg": round(price - prev, 4), "chg_pct": pct(price, prev),
            "m1": pct(price, back(30)), "ytd": pct(price, ytd_base), "asof": asof,
            "hist": [[t, round(p, 4)] for t, p in pts if t >= NOW.timestamp() - 35 * 86400]}

def atym():
    d = yahoo("ATYM.L"); return d

def fx():
    g, e = yahoo("GBPUSD=X"), yahoo("EURUSD=X")
    return {"gbpusd": g["price"], "gbpusd_chg_pct": g["chg_pct"], "eurusd": e["price"], "eurusd_chg_pct": e["chg_pct"], "asof": g["asof"]}

def westmetall(field):
    t = get(f"https://www.westmetall.com/en/markdaten.php?action=table&field={field}").text
    rows = re.findall(r"<tr[^>]*>\s*<td[^>]*>([^<]*\d{4})</td>\s*<td[^>]*>([^<]*)</td>\s*<td[^>]*>([^<]*)</td>\s*<td[^>]*>([^<]*)</td>", t)
    rows = [r for r in rows if re.search(r"\d", r[1])]
    d, cash, m3, stock = rows[0]
    num = lambda s: float(s.replace(",", "")) if re.search(r"\d", s) else None
    return {"date": dt.datetime.strptime(d.strip(), "%d. %B %Y").strftime("%Y-%m-%d"), "cash": num(cash), "m3": num(m3), "stock": num(stock),
            "prev_cash": num(rows[1][1]) if len(rows) > 1 else None}

def copper():
    hg = yahoo("HG=F")
    gbp = (OUT.get("fx") or {}).get("gbpusd") or (OLD.get("fx") or {}).get("gbpusd")
    out = {"comex_lb": hg["price"], "comex_t": round(hg["price"] * LB_PER_T, 1), "comex_gbp_t": round(hg["price"] * LB_PER_T / gbp, 1) if gbp else None,
           "comex_chg_pct": hg["chg_pct"], "comex_m1": hg["m1"], "comex_ytd": hg["ytd"], "asof": hg["asof"], "hist": hg["hist"]}
    try:
        w = westmetall("LME_Cu_cash")
        out.update(lme_cash=w["cash"], lme_3m=w["m3"], lme_date=w["date"], lme_prev=w["prev_cash"],
                   spread_t=round(out["comex_t"] - w["cash"], 1), lme_stock=w["stock"])
    except Exception as e:
        OUT["errors"]["lme"] = str(e)[:200]
        for k in ("lme_cash", "lme_3m", "lme_date", "lme_prev", "lme_stock"):
            if k in (OLD.get("copper") or {}): out[k] = OLD["copper"][k]
        if out.get("lme_cash"): out["spread_t"] = round(out["comex_t"] - out["lme_cash"], 1)
    return out

def gas():
    import openpyxl
    r = get(f"https://www.mibgas.es/en/file-access/MIBGAS_Data_{NOW.year}.xlsx?path=AGNO_{NOW.year}/XLS")
    ws = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True, data_only=True)["Trading Data PVB&VTP"]
    rows = [x for x in ws.iter_rows(min_row=2, values_only=True) if x[1] == "GDAES_D+1" and x[2] == "PVB" and (x[6] or x[7])]
    rows.sort(key=lambda x: x[0], reverse=True)
    p = lambda x: x[6] if x[6] is not None else x[7]
    return {"price": p(rows[0]), "delivery": rows[0][4].strftime("%Y-%m-%d"), "trading_day": rows[0][0].strftime("%Y-%m-%d"),
            "prev": p(rows[1]) if len(rows) > 1 else None, "asof": rows[0][0].strftime("%Y-%m-%dT12:00:00Z"),
            "hist": [[x[4].strftime("%Y-%m-%d"), p(x)] for x in rows[:30]][::-1]}

def power():
    s = (NOW - dt.timedelta(days=1)).strftime("%Y-%m-%dT00:00"); e = (NOW + dt.timedelta(days=1)).strftime("%Y-%m-%dT23:59")
    j = get(f"https://apidatos.ree.es/en/datos/mercados/precios-mercados-tiempo-real?start_date={s}&end_date={e}&time_trunc=hour").json()
    spot = next(i for i in j["included"] if i["id"] == "600" or "spot" in i["attributes"]["title"].lower())
    vals = spot["attributes"]["values"]
    days = {}
    for v in vals: days.setdefault(v["datetime"][:10], []).append(v["value"])
    last = sorted(days)[-1]; prevd = sorted(days)[-2] if len(days) > 1 else None
    avg = lambda d: round(sum(days[d]) / len(days[d]), 2)
    return {"day": last, "avg": avg(last), "min": min(days[last]), "max": max(days[last]), "prev_avg": avg(prevd) if prevd else None,
            "asof": ISO()}

def ttf():
    d = yahoo("TTF=F", "3mo"); d.pop("hist"); return d

def diesel():
    import openpyxl
    pg = get("https://energy.ec.europa.eu/data-and-analysis/weekly-oil-bulletin_en").text
    href = re.search(r'href="(/document/download/[^"]+prices%20with%20Taxes[^"]*)"', pg).group(1)
    ws = openpyxl.load_workbook(io.BytesIO(get("https://energy.ec.europa.eu" + href).content), read_only=True, data_only=True).worksheets[0]
    rows = list(ws.iter_rows(max_row=60, values_only=True))
    date = rows[1][0]; es = next(r for r in rows if str(r[0]).strip() == "Spain")
    return {"price": round(es[2], 1), "date": date.strftime("%Y-%m-%d"), "asof": date.strftime("%Y-%m-%dT12:00:00Z")}

def inventories():
    out = {}
    try:
        w = westmetall("LME_Cu_cash"); out["lme"] = {"t": w["stock"], "date": w["date"]}
    except Exception as e: OUT["errors"]["inv_lme"] = str(e)[:200]
    try:
        import xlrd
        b = get("https://www.cmegroup.com/delivery_reports/Copper_Stocks.xls").content
        sh = xlrd.open_workbook(file_contents=b).sheet_by_index(0)
        for i in range(sh.nrows):
            row = sh.row_values(i)
            if str(row[0]).strip().upper().startswith("TOTAL COPPER"):
                st = [c for c in row if isinstance(c, float)]
                out["comex"] = {"st": st[-1], "t": round(st[-1] * 0.90718474), "date": NOW.strftime("%Y-%m-%d")}; break
    except Exception as e: OUT["errors"]["inv_comex"] = str(e)[:200]
    old = OLD.get("inventories") or {}
    for k in ("lme", "comex", "shfe"):
        if k not in out and k in old: out[k] = dict(old[k], stale=True)
    if not out: raise RuntimeError("no inventory sources")
    return out

PEERS = [("FCX", "Freeport-McMoRan"), ("ANTO.L", "Antofagasta"), ("SCCO", "Southern Copper"), ("FM.TO", "First Quantum"),
         ("LUN.TO", "Lundin Mining"), ("CS.TO", "Capstone Copper"), ("HBM", "Hudbay"), ("CAML.L", "Central Asia Metals"), ("SFR.AX", "Sandfire")]
def peers():
    rows, oldrows = [], {r["sym"]: r for r in (OLD.get("peers") or {}).get("rows", [])}
    for s, n in PEERS:
        try:
            d = yahoo(s, "1y"); d.pop("hist"); rows.append(dict(d, sym=s, name=n))
        except Exception:
            if s in oldrows: rows.append(dict(oldrows[s], stale=True))
            else: rows.append({"sym": s, "name": n, "unavailable": True})
    return {"rows": rows}

FEEDS = ['"Atalaya Mining"', 'Riotinto copper', '"Proyecto Touro"', '"Masa Valverde"', '"E-LIX"', '"copper price"']
def news():
    items = []
    urls = [f"https://news.google.com/rss/search?q={urllib.parse.quote(q + ' when:30d')}&hl=en-GB&gl=GB&ceid=GB:en" for q in FEEDS]
    urls.append("https://www.mining.com/feed/")
    for u in urls:
        try:
            x = get(u).text
            for it in re.findall(r"<item>(.*?)</item>", x, re.S)[:20]:
                g = lambda tag: html.unescape(re.sub(r"<!\[CDATA\[|\]\]>", "", (re.search(f"<{tag}[^>]*>(.*?)</{tag}>", it, re.S) or [None, ""])[1])).strip()
                title, link, pub, src = g("title"), g("link"), g("pubDate"), g("source")
                if "mining.com" in u and not re.search(r"copper", title, re.I): continue
                try: ts = email.utils.parsedate_to_datetime(pub).timestamp()
                except Exception: ts = 0
                if src and title.endswith(" - " + src): title = title[: -len(src) - 3]
                items.append({"title": title, "link": link, "ts": ts, "source": src or ("MINING.COM" if "mining.com" in u else "")})
        except Exception as e: OUT["errors"]["news:" + u[:60]] = str(e)[:120]
    seen, out = set(), []
    for i in sorted(items, key=lambda i: -i["ts"]):
        k = re.sub(r"\W+", "", i["title"].lower())[:60]
        if k in seen: continue
        seen.add(k); out.append(i)
    if not out: raise RuntimeError("no news")
    return {"items": out[:15]}

for name, fn in [("atym", atym), ("fx", fx), ("copper", copper), ("gas", gas), ("power", power), ("ttf", ttf), ("diesel", diesel),
                 ("inventories", inventories), ("peers", peers), ("news", news)]:
    section(name, fn)

# relative performance ATYM vs copper
try:
    a, c = OUT["atym"], OUT["copper"]
    OUT["atym"]["rel_m1"] = round(a["m1"] - c["comex_m1"], 2); OUT["atym"]["rel_ytd"] = round(a["ytd"] - c["comex_ytd"], 2)
except Exception: pass

json.dump(OUT, open("data.json", "w"), separators=(",", ":"))
print(json.dumps(OUT["errors"], indent=1))
